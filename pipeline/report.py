"""Word 报告生成 — 输入 dataset.json + metrics.json + narratives/*.yaml +
*.png，直接产出 docx。不再从 markdown 反向解析。"""

from __future__ import annotations

import json
import os
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x21, 0x8C, 0x4E)
GREY = RGBColor(0x66, 0x66, 0x66)


def _set_cjk(doc):
    st = doc.styles["Normal"]
    st.font.name = "PingFang SC"
    st.font.size = Pt(10.5)
    try:
        from docx.oxml.ns import qn
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    except Exception:
        pass


def _add_kv(table, k, v, vc=None):
    c = table.add_row().cells
    c[0].text = k
    r = c[1].paragraphs[0].add_run(v)
    if vc:
        r.font.color.rgb = vc
    return c


def run(dataset_path: str, metrics_path: str, narr_dir: str, out_dir: str):
    import yaml
    ds = json.load(open(dataset_path))
    mt = json.load(open(metrics_path))
    syms = list(ds["matches"].keys())

    doc = Document()
    _set_cjk(doc)

    # 封面
    doc.add_heading("币安永续「资金费率 1h + 负费率」事件分析报告", level=0)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("窗口 2026-05-01→2026-05-31 · 数据源 data.binance.vision")
    r.font.size = Pt(10); r.font.color.rgb = GREY
    doc.add_paragraph()

    # 一、执行摘要
    doc.add_heading("一、执行摘要", level=1)
    doc.add_paragraph(
        "在 2026 年 5 月，币安 USDⓈ-M 永续合约中共有 11 个标的，资金费率结算周期"
        "被临时缩短为 1 小时、且区间内出现负费率。「结算周期变 1h」是币安对剧烈波动"
        "的风控信号；「负费率」意味着永续相对现货贴水。两者叠加标志这些币当时正处于"
        "极端行情。"
    )
    doc.add_paragraph(
        "关键结论：高负费率本身不等于做多机会。把 1h 区间退出后 14 天的实际价格走势"
        "拿出来看，11 个币里有 7 个继续大幅下跌（最深 IRYS -34%、PROVE -31%），靠"
        "做多吃负费率的收益远远覆盖不了本金下跌——这类是典型的「费率诱饵/接飞刀」。"
        "真正确定性较好的机会集中在价格中性的结构性品种（股票永续）。"
    )

    # 二、机会评估
    doc.add_heading("二、确定性机会评估（基于实际价格）", level=1)
    cols = ["币种", "1h内价格", "退出后14天最低", "退出后14天末值",
            "累计负费率", "verdict"]
    tb = doc.add_table(rows=1, cols=6); tb.style = "Light Grid Accent 1"
    for i, c in enumerate(cols):
        tb.rows[0].cells[i].paragraphs[0].add_run(c).bold = True
    for sym in syms:
        if sym not in mt:
            continue
        m = mt[sym]["regimes"][0]
        cells = tb.add_row().cells
        vals = [sym.replace("USDT", ""),
                f"{m['chg_in_1h_pct']:+.1f}%",
                f"{m['post_min_chg']:+.1f}%",
                f"{m['post_last_chg']:+.1f}%",
                f"{m['cum_neg_pct']:.1f}%",
                ds.get("verdict", "")]
        for i, v in enumerate(vals):
            run = cells[i].paragraphs[0].add_run(v)
            if i == 3:
                run.font.color.rgb = GREEN if m['post_last_chg'] >= 0 else RED
    doc.add_paragraph()
    # 结论三段（从 narratives 的 verdict 聚合）
    good = [s for s in syms if _get_verdict(narr_dir, s) == "✓ 较好"]
    trap = [s for s in syms if _get_verdict(narr_dir, s) == "✗ 接飞刀"]
    neutral = [s for s in syms if _get_verdict(narr_dir, s) not in ("✓ 较好", "✗ 接飞刀")]
    doc.add_heading("结论总结", level=2)
    if good:
        doc.add_paragraph(f"✓ 确定性较好: {', '.join(s.replace('USDT','') for s in good)}")
    if neutral:
        doc.add_paragraph(f"△ 中性: {', '.join(s.replace('USDT','') for s in neutral)}")
    if trap:
        doc.add_paragraph(f"✗ 接飞刀: {', '.join(s.replace('USDT','') for s in trap)}")

    # 三、方法论
    doc.add_heading("三、方法论与数据源", level=1)
    doc.add_paragraph("数据全部来自币安公开转储 data.binance.vision:")
    for b in [
        "S3 XML 索引枚举全部 772 个曾有 funding 文件的标的（含已下架）。",
        "月度 fundingRate dump 含 funding_interval_hours 列，interval==1 即 1h。",
        "1h klines dump 的 OHLC。定价区间 ≥3 次连续 1h 结算且含负费率。",
        "采样：1h 区间前 1 周每 4h、1h 内每 1h、恢复后每 4h 至 14 天。",
    ]:
        doc.add_paragraph(b, style="List Bullet")

    # 四、机制表
    doc.add_heading("四、三类驱动机制", level=1)
    tb = doc.add_table(rows=1, cols=3); tb.style = "Light Grid Accent 1"
    for i, c in enumerate(["类型", "机制", "代表币种"]):
        tb.rows[0].cells[i].paragraphs[0].add_run(c).bold = True
    for typ, mech, sms in [
        ("A. 暴涨逼空", "现货被买爆→永续被拥挤做空→负费率→空头踩踏",
         "IRYS/ID/HIVE/ORCA/PROVE/PROMPT/PRL"),
        ("B. 暴雷砸盘", "解锁/私钥泄露→现货瞬间崩盘→恐慌做空→负费率封顶",
         "GUA/LAYER"),
        ("C. 股票永续休市", "美股休市时指数冻结→预期跳空压低永续→折价负费率",
         "AAPL/MSFT"),
    ]:
        c = tb.add_row().cells
        c[0].text = typ; c[1].text = mech; c[2].text = sms

    # 五、每币详情
    doc.add_heading("五、各币详情", level=1)
    for sym in syms:
        if sym not in mt:
            continue
        doc.add_page_break()
        narr = _load_narr(narr_dir, sym)
        m = mt[sym]["regimes"][0]
        doc.add_heading(f"{sym} · {narr.get('label','?')}", level=1)

        # 指标小表
        tb = doc.add_table(rows=0, cols=2); tb.style = "Light List Accent 1"
        _add_kv(tb, "1h 区间 (UTC)", f"{m['start_str'][:16]} → {m['end_str'][:16]}")
        _add_kv(tb, "结算次数 / 负次数", f"{m.get('n','?')} / {m.get('n_neg','?')}")
        _add_kv(tb, "最负单次费率", f"{m['min_rate']*100:.3f}%", RED)
        _add_kv(tb, "累计负费率", f"{m['cum_neg_pct']:.2f}%", RED)
        _add_kv(tb, "区间内价格变化", f"{m.get('chg_in_1h_pct',''):+.1f}%",
                GREEN if m.get('chg_in_1h_pct', 0) >= 0 else RED)
        _add_kv(tb, "退出后14天最低/最高",
                f"{m.get('post_min_chg',''):+.1f}% / {m.get('post_max_chg',''):+.1f}%")
        _add_kv(tb, "退出后14天末值", f"{m.get('post_last_chg',''):+.1f}%",
                GREEN if m.get('post_last_chg', 0) >= 0 else RED)
        _add_kv(tb, "Verdict", narr.get("verdict", "—"),
                GREEN if narr.get("verdict","").startswith("✓") else None)

        if narr.get("body"):
            doc.add_heading("叙事原因", level=3)
            doc.add_paragraph(narr["body"])

        # 图
        png = os.path.join(out_dir, f"{sym}.png")
        if os.path.exists(png):
            doc.add_picture(png, width=Inches(6.4))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 完整数据表（pre + 1h + post，全量，不截断）
        doc.add_heading("完整采样数据", level=3)
        # 从 dataset raw 重建三阶段表格
        reg = ds["matches"][sym]["regimes"][0]
        from core.timewin import dt_from_ms, ms, three_phase, HOUR
        from datetime import timedelta
        start_dt = dt_from_ms(reg["start_ms"])
        end_dt = dt_from_ms(reg["end_ms"])
        pre_days = ds["config"]["window"]["pre_lookback_days"]
        post_days = ds["config"]["window"]["post_horizon_days"]
        pre0, _, _, post1 = three_phase(start_dt, end_dt, pre_days, post_days)
        price_map = {k["t"]: k["c"] for k in ds["matches"][sym]["raw"].get("price", [])}
        fund_lookup = {r["t"]: r["rate"] for r in ds["matches"][sym]["raw"].get("funding", [])}

        rows = []
        t = pre0
        while t <= start_dt - timedelta(hours=4):
            tms = ms(t)
            p = price_at(price_map, tms)
            fr = _nearest(fund_lookup, tms, 2 * 60 * 60 * 1000)
            rows.append(("pre (4h)", t, p, fr))
            t += timedelta(hours=4)
        t = start_dt
        while t <= end_dt:
            tms = ms(t)
            p = price_at(price_map, tms)
            fr = _nearest(fund_lookup, tms, 35 * 60 * 1000)
            rows.append(("1h regime", t, p, fr))
            t += timedelta(hours=1)
        t = end_dt + timedelta(hours=4)
        while t <= post1:
            tms = ms(t)
            p = price_at(price_map, tms)
            fr = _nearest(fund_lookup, tms, 2 * 60 * 60 * 1000)
            rows.append(("post (4h)", t, p, fr))
            t += timedelta(hours=4)

        tb = doc.add_table(rows=1, cols=4); tb.style = "Light Grid Accent 1"
        for ci, cn in enumerate(["阶段", "时间 (UTC)", "币价", "资金费率"]):
            tb.rows[0].cells[ci].paragraphs[0].add_run(cn).bold = True
        for phase, dt, pri, rate in rows:
            cl = tb.add_row().cells
            cl[0].text = phase
            cl[1].text = f"{dt:%Y-%m-%d %H:%M}"
            cl[2].text = _fmt_p(pri) if pri else "—"
            rv = f"{rate*100:+.4f}%" if rate else ""
            rc = cl[3].paragraphs[0].add_run(rv)
            if rate is not None and rate < 0:
                rc.font.color.rgb = RED
            elif rate is not None and rate > 0:
                rc.font.color.rgb = GREEN
        doc.add_paragraph(f"（共 {len(rows)} 行采样，已全部列出）",
                          style="Intense Quote")

        # 来源
        if narr.get("sources"):
            doc.add_heading("参考来源", level=3)
            for u in narr["sources"]:
                doc.add_paragraph(u, style="List Bullet")

    path = os.path.join(out_dir, "资金费率1h负费率事件分析报告.docx")
    doc.save(path)
    print(f"saved {path}", file=sys.stderr)


def _load_narr(narr_dir, sym):
    import yaml
    p = os.path.join(narr_dir, f"{sym}.yaml")
    return yaml.safe_load(open(p)) if os.path.exists(p) else {}


def _get_verdict(narr_dir, sym):
    return _load_narr(narr_dir, sym).get("verdict", "")


def price_at(pm, tms):
    best = None
    for pt in sorted(pm):
        if pt <= tms:
            best = pm[pt]
        else:
            break
    return best if best is not None else (pm[min(pm.keys())] if pm else None)


def _nearest(fund, tms, tol):
    best, bestd = None, None
    lo, hi = tms - tol, tms + tol
    for tt in fund:
        if lo <= tt <= hi:
            d = abs(tt - tms)
            if bestd is None or d < bestd:
                best, bestd = fund[tt], d
    return best


def _fmt_p(p):
    if p is None:
        return "—"
    if p >= 1:
        return f"{p:,.4f}"
    return f"{p:.8f}".rstrip("0").rstrip(".")
